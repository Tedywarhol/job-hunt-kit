#!/usr/bin/env python3
"""Extrait le texte brut d'un CV ou document existant (PDF, DOCX, TXT, MD).

Brique déterministe de l'import de profil : elle ne fait qu'extraire le texte, elle ne
l'interprète pas. L'interprétation (quelle expérience, quel projet, quelles compétences
en tirer, comment les fusionner dans templates/cv/cv-data.json sans écraser ni inventer)
est un travail de compréhension du langage : elle est faite par le skill Claude Code
`import-profile` (.claude/skills/import-profile/SKILL.md), jamais par un parseur regex —
un CV a un format trop libre pour qu'une extraction structurelle déterministe soit fiable
sans risquer d'écrire une donnée fausse sur le profil réel du candidat.

Usage:
  python scripts/extract_document_text.py <chemin/vers/cv.pdf>
  python scripts/extract_document_text.py <chemin/vers/cv.docx> --out state/import-source.txt
"""
import argparse
import os
import sys
from typing import Optional

ROOT: str = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from logutil import log_error

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")


def extract_pdf(path: str) -> str:
    import pypdf
    reader = pypdf.PdfReader(path)
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as e:
            log_error(f"extract_document_text: page PDF illisible dans {path}", e)
    return "\n".join(pages)


def extract_docx(path: str) -> str:
    import docx
    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def extract_text_file(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def extract_text(path: str) -> str:
    """Extrait le texte brut de `path` selon son extension. Lève ValueError si non supporté."""
    ext = os.path.splitext(path)[1].lower()
    if not os.path.isfile(path):
        raise FileNotFoundError(f"Fichier introuvable : {path}")
    if ext == ".pdf":
        text = extract_pdf(path)
    elif ext == ".docx":
        text = extract_docx(path)
    elif ext in (".txt", ".md"):
        text = extract_text_file(path)
    else:
        raise ValueError(
            f"Extension non supportée : '{ext}'. Formats acceptés : {', '.join(SUPPORTED_EXTENSIONS)}."
        )
    text = text.strip()
    if not text:
        raise ValueError(f"Aucun texte extrait de {path} — fichier vide, scanné en image, ou protégé.")
    return text


def main() -> None:
    ap = argparse.ArgumentParser(description="Extrait le texte brut d'un CV/document (PDF, DOCX, TXT, MD).")
    ap.add_argument("path", help="Chemin du fichier source")
    ap.add_argument("--out", help="Fichier de sortie (texte). Défaut : affiche sur stdout.")
    args = ap.parse_args()

    try:
        text = extract_text(args.path)
    except (FileNotFoundError, ValueError) as e:
        sys.exit(str(e))

    if args.out:
        os.makedirs(os.path.dirname(os.path.abspath(args.out)) or ".", exist_ok=True)
        with open(args.out, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"Texte extrait ({len(text)} caractères) -> {args.out}")
    else:
        print(text)


if __name__ == "__main__":
    main()
